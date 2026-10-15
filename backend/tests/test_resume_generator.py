"""The resume must contain only the requesting user's own data."""
from docx import Document
from services.resume_generator import _generate_with_multiplier, contact_line


def _text(path):
    return "\n".join(p.text for p in Document(path).paragraphs) + "\n".join(
        c.text for t in Document(path).tables for r in t.rows for c in r.cells
    )


def test_resume_uses_only_the_given_users_education_and_skills(tmp_path):
    out = tmp_path / "r.docx"
    _generate_with_multiplier(
        profile={"name": "Ada Lovelace", "email": "ada@example.com", "phone": None, "linkedin": None, "github": None},
        skills={"Languages": "Rust, OCaml"},
        experience=[],
        projects=[],
        education=[{"degree": "B.Sc. Mathematics", "school": "University of London", "track": None, "relevant_courses": "Analysis"}],
        output_path=str(out),
    )
    text = _text(out)
    assert "B.Sc. Mathematics" in text and "University of London" in text and "Relevant: Analysis" in text
    assert "Rust, OCaml" in text
    for owner_only in ("UT Dallas", "Sarvajanik", "CometX", "Vishwaa", "Python, JavaScript"):
        assert owner_only not in text


def test_contact_line_skips_missing_fields():
    assert contact_line({"phone": None, "email": "a@b.co", "linkedin": "", "github": "gh/a"}) == "a@b.co  |  gh/a"
